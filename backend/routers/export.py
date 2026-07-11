from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import sys, os, io
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from database import get_db
from models import Investment, CreditCard, BankAccount, NetWorthSnapshot
from datetime import datetime

from routers.auth import require_auth

router = APIRouter(prefix="/export", tags=["export"], dependencies=[Depends(require_auth)])


@router.get("/excel")
def export_excel(db: Session = Depends(get_db)):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    hfill = PatternFill(start_color="0d1320", end_color="0d1320", fill_type="solid")
    hfont = Font(color="00d4aa", bold=True)

    def make_sheet(title, headers, rows):
        ws = wb.create_sheet(title=title)
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = hfont; cell.fill = hfill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(ci)].width = 22
        for ri, row in enumerate(rows, 2):
            for ci, val in enumerate(row, 1):
                ws.cell(row=ri, column=ci, value=val)

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    investments = db.query(Investment).all()
    make_sheet("Investments",
        ["Name", "Asset Class", "Platform", "Invested (₹)", "Current Value (₹)", "P&L (₹)", "Return %", "Last Updated"],
        [(i.name, i.asset_class, i.platform, i.invested_amount, i.current_value,
          round(i.current_value - i.invested_amount, 2),
          round((i.current_value - i.invested_amount) / i.invested_amount * 100, 2) if i.invested_amount else 0,
          i.last_updated.strftime("%Y-%m-%d")) for i in investments])

    cards = db.query(CreditCard).all()
    make_sheet("Credit Cards",
        ["Name", "Bank", "Credit Limit (₹)", "Balance (₹)", "Utilization %", "Due Day"],
        [(c.name, c.bank, c.credit_limit, c.current_balance,
          round(c.current_balance / c.credit_limit * 100, 1) if c.credit_limit else 0,
          c.due_day) for c in cards])

    accounts = db.query(BankAccount).all()
    make_sheet("Bank Accounts",
        ["Name", "Bank", "Purpose", "Balance (₹)", "Monthly Inflow", "Monthly Outflow", "Emergency Fund"],
        [(a.name, a.bank, a.purpose, a.balance, a.monthly_inflow, a.monthly_outflow,
          "Yes" if a.is_emergency_fund else "No") for a in accounts])

    snapshots = db.query(NetWorthSnapshot).order_by(NetWorthSnapshot.snapshot_date).all()
    make_sheet("Net Worth History",
        ["Date", "Net Worth (₹)", "Total Assets (₹)", "Liabilities (₹)", "Investments (₹)", "Bank Balance (₹)"],
        [(s.snapshot_date.strftime("%Y-%m-%d"), s.net_worth, s.total_assets,
          s.total_liabilities, s.investment_value, s.bank_balance) for s in snapshots])

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    fname = f"wealthos_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"})


@router.get("/backup-db")
def backup_db():
    from database import DB_PATH
    db_path = DB_PATH
    if not os.path.exists(db_path):
        from fastapi import HTTPException
        raise HTTPException(status_code=404, detail="Database not found")
    with open(db_path, "rb") as f:
        data = f.read()
    fname = f"wealthos_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
    return StreamingResponse(io.BytesIO(data), media_type="application/octet-stream",
        headers={"Content-Disposition": f"attachment; filename={fname}"})
